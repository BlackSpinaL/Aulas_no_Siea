import io
import pandas as pd
import streamlit as st

# Tentar importar openpyxl para formatação avançada do Excel
try:
    import openpyxl
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter
    HAS_OPENPYXL = True
except ImportError:
    HAS_OPENPYXL = False

st.set_page_config(
    page_title="Conferência dos Diários",
    layout="wide",
)

# Estilização CSS personalizada para a Tabela estilo Excel
st.markdown(
    """
<style>
.excel-table {
    width: 100%;
    border-collapse: collapse;
    font-family: 'Segoe UI', Tahoma, Geneva, Verdana, sans-serif;
    font-size: 13px;
}
.excel-table th {
    background-color: #1b55a8;
    color: white;
    font-weight: bold;
    text-align: center;
    border: 1px solid #103871;
    padding: 5px;
}
.excel-table .header-main {
    background-color: #0d2b59;
    color: white;
    font-size: 16px;
    font-weight: bold;
    text-align: center;
    padding: 8px;
}
.excel-table td {
    border: 1px solid #cccccc;
    padding: 6px;
    text-align: center;
    color: #333333;
}
.row-materia-1 {
    background-color: #e8f5e9 !important; /* Verde bem claro */
}
.row-materia-0 {
    background-color: #ffffff !important; /* Branco */
}
.aula-destaque {
    font-weight: 900 !important;
    font-size: 14px !important;
    color: #000000 !important;
    background-color: #fff59d !important; /* Amarelo destaque */
}
.status-ok {
    background-color: #c8e6c9 !important;
    color: #2e7d32 !important;
    font-weight: bold;
}
.status-err {
    background-color: #ffcdd2 !important;
    color: #c62828 !important;
    font-weight: bold;
}
.action-table {
    width: 100%;
    border-collapse: collapse;
    margin-bottom: 10px;
}
.action-table th {
    background-color: #0d2b59;
    color: white;
    padding: 6px;
    font-size: 13px;
    text-align: left;
}
.action-table td {
    border: 1px solid #e0e0e0;
    padding: 8px;
    background-color: #f9f9f9;
}
</style>
""",
    unsafe_allow_html=True,
)

# Inicialização do Session State
if "ano_letivo" not in st.session_state:
    st.session_state.ano_letivo = 2026
if "turmas" not in st.session_state:
    st.session_state.turmas = ["12101", "12102", "12103", "12104"]
if "turma_atual" not in st.session_state:
    st.session_state.turma_atual = "12101"
if "editing_idx" not in st.session_state:
    st.session_state.editing_idx = None

DISCIPLINAS_PADRAO = [
    "ARTE",
    "BIOLOGIA",
    "BIOLOGIA NA PRATICA",
    "C. DA NATUREZA P/ ENEM",
    "CIENCIAS",
    "DESENV. SUSTENTAVEL",
    "ED. FISICA NA PRATICA",
    "ED. PARA PROFISSOES",
    "ED. SOCIO. ENS. RELIG.",
    "EDUCACAO FINANCEIRA",
    "EDUCACAO FISICA",
    "FILOSOFIA",
    "FISICA",
    "GEOGRAFIA",
    "HISTORIA",
    "LABORATORIO / BIOLOGIA",
    "LABORATORIO / CIENCIAS",
    "LABORATORIO / FISICA",
    "LABORATORIO / QUIMICA",
    "LINGUA INGLESA",
    "LÍNGUA INGLESA NA PRATICA",
    "LINGUA PORTUGUESA",
    "LINGUA PORTUGUESA 2",
    "MAT. E ESTATISTICA",
    "MATEMATICA",
    "OFICINA DE TEXTO",
    "PROJETO DE VIDA",
    "QUIMICA",
    "QUIMICA NA PRATICA",
    "SOCIOLOGIA",
]

if "disciplinas" not in st.session_state:
    st.session_state.disciplinas = DISCIPLINAS_PADRAO.copy()

if "dias_etapas" not in st.session_state:
    st.session_state.dias_etapas = pd.DataFrame({
        "Dia da Semana": ["Segunda", "Terça", "Quarta", "Quinta", "Sexta"],
        "Etapa 1": [13, 12, 14, 14, 13],
        "Etapa 2": [15, 16, 13, 13, 14],
        "Etapa 3": [12, 12, 13, 13, 13],
    })

if "lancamentos" not in st.session_state:
    st.session_state.lancamentos = []

st.title(f"📊 Conferência dos Diários - {st.session_state.ano_letivo}")
aba1, aba2 = st.tabs(["📋 Grade de Aulas (Visualização)", "⚙️ Configurações"])

# -------------------------------------------------------------
# ABA 2: CONFIGURAÇÕES (ANO, TURMAS, MATÉRIAS E CALENDÁRIO)
# -------------------------------------------------------------
with aba2:
    st.subheader("⚙️ Configurações Gerais")
    col_cfg1, col_cfg2, col_cfg3 = st.columns([1, 1, 1])
    with col_cfg1:
        st.markdown("### 1. Ano Letivo")
        novo_ano = st.number_input(
            "Ano Letivo:",
            min_value=2020,
            max_value=2035,
            value=st.session_state.ano_letivo,
        )
        if novo_ano != st.session_state.ano_letivo:
            st.session_state.ano_letivo = novo_ano
            st.rerun()

    with col_cfg2:
        st.markdown("### 2. Gerenciar Turmas")
        nova_turma = st.text_input("Nova Turma (ex: 12105):")
        if st.button("➕ Adicionar Turma"):
            if (
                nova_turma.strip()
                and nova_turma.upper() not in st.session_state.turmas
            ):
                st.session_state.turmas.append(nova_turma.upper())
                st.session_state.turmas.sort()
                st.rerun()
        turma_rem = st.selectbox(
            "Excluir Turma Definitivamente:",
            ["-- Selecione --"] + st.session_state.turmas,
        )
        if st.button("🗑️ Deletar Turma") and turma_rem != "-- Selecione --":
            st.session_state.turmas.remove(turma_rem)
            st.session_state.lancamentos = [
                l
                for l in st.session_state.lancamentos
                if l.get("turma") != turma_rem
            ]
            if st.session_state.turmas:
                st.session_state.turma_atual = st.session_state.turmas[0]
            st.rerun()

    with col_cfg3:
        st.markdown("### 3. Gerenciar Disciplinas")
        
        # Adicionar nova disciplina
        nova_mat = st.text_input("Nova Disciplina:")
        if st.button("➕ Adicionar Disciplina"):
            if (
                nova_mat.strip()
                and nova_mat.upper() not in st.session_state.disciplinas
            ):
                st.session_state.disciplinas.append(nova_mat.upper())
                st.session_state.disciplinas.sort()
                st.rerun()
        
        # Deletar disciplina existente
        st.markdown("**Excluir Disciplina:**")
        mat_rem = st.selectbox(
            "Selecione a disciplina:",
            ["-- Selecione --"] + st.session_state.disciplinas,
            key="del_disc_select"
        )
        if st.button("🗑️ Deletar Disciplina", key="del_disc_btn"):
            if mat_rem != "-- Selecione --":
                st.session_state.disciplinas.remove(mat_rem)
                # Remove lançamentos dessa disciplina em todas as turmas
                st.session_state.lancamentos = [
                    l for l in st.session_state.lancamentos
                    if l.get("materia") != mat_rem
                ]
                st.success(f"Disciplina '{mat_rem}' excluída!")
                st.rerun()

    st.markdown("---")
    
    # -------------------------------------------------------------
    # EDIÇÃO DOS DIAS DE CADA ETAPA (NOVO RECURSO)
    # -------------------------------------------------------------
    st.markdown("### 📅 Editar Dias de Cada Etapa")
    st.caption("Altere os valores abaixo e clique em 'Salvar Alterações' para atualizar o cálculo.")
    
    with st.form("form_editar_dias"):
        df_edit = st.session_state.dias_etapas.copy()
        
        # Cria colunas para edição
        cols = st.columns([1.5, 1, 1, 1])
        with cols[0]:
            st.markdown("**Dia da Semana**")
        with cols[1]:
            st.markdown("**Etapa 1**")
        with cols[2]:
            st.markdown("**Etapa 2**")
        with cols[3]:
            st.markdown("**Etapa 3**")
        
        # Inputs para cada linha
        novos_valores = []
        for i, row in df_edit.iterrows():
            cols = st.columns([1.5, 1, 1, 1])
            with cols[0]:
                st.markdown(f"**{row['Dia da Semana']}**")
            with cols[1]:
                e1 = st.number_input(
                    f"E1_{row['Dia da Semana']}",
                    min_value=0,
                    max_value=60,
                    value=int(row['Etapa 1']),
                    label_visibility="collapsed",
                    key=f"e1_{i}"
                )
            with cols[2]:
                e2 = st.number_input(
                    f"E2_{row['Dia da Semana']}",
                    min_value=0,
                    max_value=60,
                    value=int(row['Etapa 2']),
                    label_visibility="collapsed",
                    key=f"e2_{i}"
                )
            with cols[3]:
                e3 = st.number_input(
                    f"E3_{row['Dia da Semana']}",
                    min_value=0,
                    max_value=60,
                    value=int(row['Etapa 3']),
                    label_visibility="collapsed",
                    key=f"e3_{i}"
                )
            novos_valores.append({
                "Dia da Semana": row['Dia da Semana'],
                "Etapa 1": e1,
                "Etapa 2": e2,
                "Etapa 3": e3,
            })
        
        if st.form_submit_button("💾 Salvar Alterações"):
            st.session_state.dias_etapas = pd.DataFrame(novos_valores)
            st.success("Dias das etapas atualizados com sucesso!")
            st.rerun()

    st.markdown("---")
    
    # Cálculo automático dos totais de dias das etapas
    df_calc = st.session_state.dias_etapas.copy()
    df_calc["TOTAL DIAS"] = df_calc[["Etapa 1", "Etapa 2", "Etapa 3"]].sum(
        axis=1
    )
    linha_total = pd.DataFrame({
        "Dia da Semana": ["TOTAL ETAPAS"],
        "Etapa 1": [df_calc["Etapa 1"].sum()],
        "Etapa 2": [df_calc["Etapa 2"].sum()],
        "Etapa 3": [df_calc["Etapa 3"].sum()],
        "TOTAL DIAS": [df_calc["TOTAL DIAS"].sum()],
    })
    df_conferencia = pd.concat([df_calc, linha_total], ignore_index=True)

    st.markdown("### 📋 CÁLCULO DE DIAS DE CADA ETAPA")
    st.dataframe(
        df_conferencia,
        use_container_width=True,
        hide_index=True,
    )

# -------------------------------------------------------------
# ABA 1: VISUALIZAÇÃO PRINCIPAL
# -------------------------------------------------------------
with aba1:
    c_t, c_dummy = st.columns([2, 4])
    with c_t:
        st.session_state.turma_atual = st.selectbox(
            "📍 Selecione a Turma para Conferência:", st.session_state.turmas
        )

    # SEÇÃO DE EDIÇÃO OU NOVO LANÇAMENTO
    if st.session_state.editing_idx is not None:
        st.markdown("### ✏️ Editando Disciplina Lançada")
        item_edit = st.session_state.lancamentos[st.session_state.editing_idx]
        c_m, c_prev, c_seg, c_ter, c_qua, c_qui, c_sex = st.columns(
            [2.5, 1.2, 1, 1, 1, 1, 1]
        )
        with c_m:
            mat_e = st.text_input(
                "Componente Curricular:",
                value=item_edit["materia"],
                disabled=True,
            )
        with c_prev:
            prev_e = st.number_input(
                "Aulas Anual:",
                min_value=0,
                value=item_edit["previsto"],
                step=10,
            )
        with c_seg:
            seg_e = st.number_input(
                "Segunda", min_value=0, max_value=10, value=item_edit["Seg"]
            )
        with c_ter:
            ter_e = st.number_input(
                "Terça", min_value=0, max_value=10, value=item_edit["Ter"]
            )
        with c_qua:
            qua_e = st.number_input(
                "Quarta", min_value=0, max_value=10, value=item_edit["Qua"]
            )
        with c_qui:
            qui_e = st.number_input(
                "Quinta", min_value=0, max_value=10, value=item_edit["Qui"]
            )
        with c_sex:
            sex_e = st.number_input(
                "Sexta", min_value=0, max_value=10, value=item_edit["Sex"]
            )

        col_sav1, col_sav2 = st.columns([1, 4])
        with col_sav1:
            if st.button("💾 Salvar Alterações"):
                st.session_state.lancamentos[st.session_state.editing_idx] = {
                    "turma": st.session_state.turma_atual,
                    "materia": item_edit["materia"],
                    "previsto": prev_e,
                    "Seg": seg_e,
                    "Ter": ter_e,
                    "Qua": qua_e,
                    "Qui": qui_e,
                    "Sex": sex_e,
                }
                st.session_state.editing_idx = None
                st.success("Alterações salvas com sucesso!")
                st.rerun()
        with col_sav2:
            if st.button("❌ Cancelar Edição"):
                st.session_state.editing_idx = None
                st.rerun()

    else:
        st.markdown("### ✏️ Lançar Matéria na Grade")
        c_m, c_prev, c_seg, c_ter, c_qua, c_qui, c_sex = st.columns(
            [2.5, 1.2, 1, 1, 1, 1, 1]
        )
        with c_m:
            mat_escolhida = st.selectbox(
                "Componente Curricular:",
                ["-- Selecione --"] + st.session_state.disciplinas,
            )
        with c_prev:
            aulas_anual_prevista = st.number_input(
                "Aulas Anual:", min_value=0, value=80, step=10
            )
        with c_seg:
            seg = st.number_input("Segunda", min_value=0, max_value=10, value=0)
        with c_ter:
            ter = st.number_input("Terça", min_value=0, max_value=10, value=0)
        with c_qua:
            qua = st.number_input("Quarta", min_value=0, max_value=10, value=0)
        with c_qui:
            qui = st.number_input("Quinta", min_value=0, max_value=10, value=0)
        with c_sex:
            sex = st.number_input("Sexta", min_value=0, max_value=10, value=0)

        if st.button("📥 Lançar na Grade"):
            if mat_escolhida != "-- Selecione --":
                st.session_state.lancamentos.append({
                    "turma": st.session_state.turma_atual,
                    "materia": mat_escolhida,
                    "previsto": aulas_anual_prevista,
                    "Seg": seg,
                    "Ter": ter,
                    "Qua": qua,
                    "Qui": qui,
                    "Sex": sex,
                })
                st.success(f"'{mat_escolhida}' lançada com sucesso!")
                st.rerun()

    st.markdown("---")

    lancamentos_turma = [
        (orig_idx, item)
        for orig_idx, item in enumerate(st.session_state.lancamentos)
        if item.get("turma", st.session_state.turma_atual)
        == st.session_state.turma_atual
    ]

    dias_map = st.session_state.dias_etapas.set_index("Dia da Semana")

    if lancamentos_turma:
        titulo_grade = f"TURMA: {st.session_state.turma_atual} - PREVISÃO DE AULAS - {st.session_state.ano_letivo}"
        html_code = f"""
        <table class="excel-table">
            <thead>
                <tr>
                    <th colspan="21" class="header-main">{titulo_grade}</th>
                </tr>
                <tr>
                    <th rowspan="2" style="width: 15%;">Componente Curricular</th>
                    <th rowspan="2" style="width: 4%;">Etapa</th>
                    <th colspan="3">Segunda</th>
                    <th colspan="3">Terça</th>
                    <th colspan="3">Quarta</th>
                    <th colspan="3">Quinta</th>
                    <th colspan="3">Sexta</th>
                    <th rowspan="2">Nº aulas<br>por etapa</th>
                    <th rowspan="2">TOTAL<br>AULAS</th>
                    <th rowspan="2">AULAS<br>ANUAL</th>
                    <th rowspan="2">SITUAÇÃO</th>
                </tr>
                <tr>
                    <th>Nº aulas</th><th>Nº/etapa</th><th>Total</th>
                    <th>Nº aulas</th><th>Nº/etapa</th><th>Total</th>
                    <th>Nº aulas</th><th>Nº/etapa</th><th>Total</th>
                    <th>Nº aulas</th><th>Nº/etapa</th><th>Total</th>
                    <th>Nº aulas</th><th>Nº/etapa</th><th>Total</th>
                </tr>
            </thead>
            <tbody>
        """

        for pos, (orig_idx, item) in enumerate(lancamentos_turma):
            mat = item["materia"]
            prev = item["previsto"]
            s_a, t_a, q_a, qui_a, sex_a = (
                item["Seg"],
                item["Ter"],
                item["Qua"],
                item["Qui"],
                item["Sex"],
            )

            bg_class = "row-materia-1" if pos % 2 == 0 else "row-materia-0"

            tot_anual = 0
            for e in [1, 2, 3]:
                s_d = dias_map.loc["Segunda", f"Etapa {e}"]
                t_d = dias_map.loc["Terça", f"Etapa {e}"]
                q_d = dias_map.loc["Quarta", f"Etapa {e}"]
                qui_d = dias_map.loc["Quinta", f"Etapa {e}"]
                sex_d = dias_map.loc["Sexta", f"Etapa {e}"]
                tot_anual += (
                    (s_a * s_d)
                    + (t_a * t_d)
                    + (q_a * q_d)
                    + (qui_a * qui_d)
                    + (sex_a * sex_d)
                )

            dif = tot_anual - prev
            if dif == 0:
                sit_txt, sit_class = "☑ OK", "status-ok"
            elif dif > 0:
                sit_txt, sit_class = f"⚠️ EXCESSO (+{dif})", "status-err"
            else:
                sit_txt, sit_class = f"❌ FALTA ({dif})", "status-err"

            # Bolding / Highlight para células mescladas
            s_a_h = f'<td rowspan="3" class="{bg_class} aula-destaque">{s_a}</td>' if s_a > 0 else f'<td rowspan="3" class="{bg_class}">{s_a}</td>'
            t_a_h = f'<td rowspan="3" class="{bg_class} aula-destaque">{t_a}</td>' if t_a > 0 else f'<td rowspan="3" class="{bg_class}">{t_a}</td>'
            q_a_h = f'<td rowspan="3" class="{bg_class} aula-destaque">{q_a}</td>' if q_a > 0 else f'<td rowspan="3" class="{bg_class}">{q_a}</td>'
            qui_a_h = f'<td rowspan="3" class="{bg_class} aula-destaque">{qui_a}</td>' if qui_a > 0 else f'<td rowspan="3" class="{bg_class}">{qui_a}</td>'
            sex_a_h = f'<td rowspan="3" class="{bg_class} aula-destaque">{sex_a}</td>' if sex_a > 0 else f'<td rowspan="3" class="{bg_class}">{sex_a}</td>'

            for e in [1, 2, 3]:
                s_d = dias_map.loc["Segunda", f"Etapa {e}"]
                t_d = dias_map.loc["Terça", f"Etapa {e}"]
                q_d = dias_map.loc["Quarta", f"Etapa {e}"]
                qui_d = dias_map.loc["Quinta", f"Etapa {e}"]
                sex_d = dias_map.loc["Sexta", f"Etapa {e}"]

                s_t, t_t, q_t, qui_t, sex_t = (
                    s_a * s_d,
                    t_a * t_d,
                    q_a * q_d,
                    qui_a * qui_d,
                    sex_a * sex_d,
                )
                tot_etapa = s_t + t_t + q_t + qui_t + sex_t

                html_code += f'<tr class="{bg_class}">'
                if e == 1:
                    html_code += f'<td rowspan="3" style="font-weight:bold;">{mat}</td>'

                html_code += f"<td>{e}</td>"

                if e == 1:
                    html_code += f"{s_a_h}<td>{s_d}</td><td>{s_t}</td>"
                    html_code += f"{t_a_h}<td>{t_d}</td><td>{t_t}</td>"
                    html_code += f"{q_a_h}<td>{q_d}</td><td>{q_t}</td>"
                    html_code += f"{qui_a_h}<td>{qui_d}</td><td>{qui_t}</td>"
                    html_code += f"{sex_a_h}<td>{sex_d}</td><td>{sex_t}</td>"
                else:
                    html_code += f"<td>{s_d}</td><td>{s_t}</td>"
                    html_code += f"<td>{t_d}</td><td>{t_t}</td>"
                    html_code += f"<td>{q_d}</td><td>{q_t}</td>"
                    html_code += f"<td>{qui_d}</td><td>{qui_t}</td>"
                    html_code += f"<td>{sex_d}</td><td>{sex_t}</td>"

                html_code += f'<td style="font-weight:bold;">{tot_etapa}</td>'

                if e == 1:
                    html_code += (
                        f'<td rowspan="3" style="font-weight:bold; font-size:14px;">{tot_anual}</td>'
                        f'<td rowspan="3" style="font-weight:bold; font-size:14px;">{prev}</td>'
                        f'<td rowspan="3" class="{sit_class}">{sit_txt}</td>'
                    )

                html_code += "</tr>"

        html_code += "</tbody></table>"
        st.markdown(html_code, unsafe_allow_html=True)

        # -------------------------------------------------------------
        # AÇÕES PARA CADA DISCIPLINA ORGANIZADAS EM TABELA
        # -------------------------------------------------------------
        st.markdown("### ⚙️ Ações para cada Disciplina")

        cols_per_row = 3
        for i in range(0, len(lancamentos_turma), cols_per_row):
            batch = lancamentos_turma[i : i + cols_per_row]
            cols = st.columns(cols_per_row)
            for j, (orig_idx, item) in enumerate(batch):
                with cols[j]:
                    st.markdown(
                        f"""
                        <table class="action-table">
                            <thead>
                                <tr><th>📘 {item['materia']}</th></tr>
                            </thead>
                        </table>
                        """,
                        unsafe_allow_html=True,
                    )
                    c_b1, c_b2 = st.columns(2)
                    with c_b1:
                        if st.button("✏️ Editar", key=f"btn_edit_{orig_idx}"):
                            st.session_state.editing_idx = orig_idx
                            st.rerun()
                    with c_b2:
                        if st.button("🗑️ Deletar", key=f"btn_del_{orig_idx}"):
                            st.session_state.lancamentos.pop(orig_idx)
                            st.rerun()

        # -------------------------------------------------------------
        # EXPORTAÇÃO EXCEL IDENTICA À TELA (OPENPYXL)
        # -------------------------------------------------------------
        st.markdown("---")
        if HAS_OPENPYXL:
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = "Conferência"
            ws.views.sheetView[0].showGridLines = True

            # Estilos
            header_main_fill = PatternFill(start_color="0D2B59", end_color="0D2B59", fill_type="solid")
            header_sub_fill = PatternFill(start_color="1B55A8", end_color="1B55A8", fill_type="solid")
            materia_bg_0 = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")
            materia_bg_1 = PatternFill(start_color="E8F5E9", end_color="E8F5E9", fill_type="solid")
            destaque_fill = PatternFill(start_color="FFF59D", end_color="FFF59D", fill_type="solid")
            ok_fill = PatternFill(start_color="C8E6C9", end_color="C8E6C9", fill_type="solid")
            err_fill = PatternFill(start_color="FFCDD2", end_color="FFCDD2", fill_type="solid")

            font_header = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
            font_title = Font(name="Segoe UI", size=13, bold=True, color="FFFFFF")
            font_regular = Font(name="Segoe UI", size=10)
            font_bold = Font(name="Segoe UI", size=10, bold=True)
            font_ok = Font(name="Segoe UI", size=10, bold=True, color="2E7D32")
            font_err = Font(name="Segoe UI", size=10, bold=True, color="C62828")

            align_center = Alignment(horizontal="center", vertical="center", wrap_text=True)
            thin_border_side = Side(style="thin", color="CCCCCC")
            header_border_side = Side(style="thin", color="103871")
            border_cell = Border(left=thin_border_side, right=thin_border_side, top=thin_border_side, bottom=thin_border_side)
            border_header = Border(left=header_border_side, right=header_border_side, top=header_border_side, bottom=header_border_side)

            # Linha 1: Título Principal
            ws.merge_cells("A1:U1")
            ws["A1"] = titulo_grade
            ws["A1"].font = font_title
            ws["A1"].fill = header_main_fill
            ws["A1"].alignment = align_center

            # Ajuste de merge do cabeçalho
            ws.merge_cells("A2:A3")
            ws["A2"] = "Componente Curricular"
            ws.merge_cells("B2:B3")
            ws["B2"] = "Etapa"
            ws.merge_cells("C2:E2")
            ws["C2"] = "Segunda"
            ws.merge_cells("F2:H2")
            ws["F2"] = "Terça"
            ws.merge_cells("I2:K2")
            ws["I2"] = "Quarta"
            ws.merge_cells("L2:N2")
            ws["L2"] = "Quinta"
            ws.merge_cells("O2:Q2")
            ws["O2"] = "Sexta"
            ws.merge_cells("R2:R3")
            ws["R2"] = "Nº aulas por etapa"
            ws.merge_cells("S2:S3")
            ws["S2"] = "TOTAL AULAS"
            ws.merge_cells("T2:T3")
            ws["T2"] = "AULAS ANUAL"
            ws.merge_cells("U2:U3")
            ws["U2"] = "SITUAÇÃO"

            sub_headers = ["Nº aulas", "Nº/etapa", "Total"] * 5
            for col_idx, sub in enumerate(sub_headers, start=3):
                cell = ws.cell(row=3, column=col_idx)
                cell.value = sub

            # Estilizar todo o bloco de cabeçalho
            for r in range(2, 4):
                for c in range(1, 22):
                    cell = ws.cell(row=r, column=c)
                    cell.fill = header_sub_fill
                    cell.font = font_header
                    cell.alignment = align_center
                    cell.border = border_header

            current_row = 4
            for pos, (orig_idx, item) in enumerate(lancamentos_turma):
                mat = item["materia"]
                prev = item["previsto"]
                s_a, t_a, q_a, qui_a, sex_a = item["Seg"], item["Ter"], item["Qua"], item["Qui"], item["Sex"]
                bg_fill = materia_bg_1 if pos % 2 == 0 else materia_bg_0

                tot_anual = 0
                for e in [1, 2, 3]:
                    s_d = dias_map.loc["Segunda", f"Etapa {e}"]
                    t_d = dias_map.loc["Terça", f"Etapa {e}"]
                    q_d = dias_map.loc["Quarta", f"Etapa {e}"]
                    qui_d = dias_map.loc["Quinta", f"Etapa {e}"]
                    sex_d = dias_map.loc["Sexta", f"Etapa {e}"]
                    tot_anual += (s_a * s_d) + (t_a * t_d) + (q_a * q_d) + (qui_a * qui_d) + (sex_a * sex_d)

                dif = tot_anual - prev
                sit_txt = "OK" if dif == 0 else (f"EXCESSO (+{dif})" if dif > 0 else f"FALTA ({dif})")
                sit_font = font_ok if dif == 0 else font_err
                sit_fill = ok_fill if dif == 0 else err_fill

                # Preenchimento das 3 linhas do componente
                r_start = current_row
                r_end = current_row + 2

                ws.merge_cells(start_row=r_start, start_column=1, end_row=r_end, end_column=1)
                ws.cell(row=r_start, column=1, value=mat)

                ws.merge_cells(start_row=r_start, start_column=3, end_row=r_end, end_column=3)
                ws.cell(row=r_start, column=3, value=s_a)

                ws.merge_cells(start_row=r_start, start_column=6, end_row=r_end, end_column=6)
                ws.cell(row=r_start, column=6, value=t_a)

                ws.merge_cells(start_row=r_start, start_column=9, end_row=r_end, end_column=9)
                ws.cell(row=r_start, column=9, value=q_a)

                ws.merge_cells(start_row=r_start, start_column=12, end_row=r_end, end_column=12)
                ws.cell(row=r_start, column=12, value=qui_a)

                ws.merge_cells(start_row=r_start, start_column=15, end_row=r_end, end_column=15)
                ws.cell(row=r_start, column=15, value=sex_a)

                ws.merge_cells(start_row=r_start, start_column=19, end_row=r_end, end_column=19)
                ws.cell(row=r_start, column=19, value=tot_anual)

                ws.merge_cells(start_row=r_start, start_column=20, end_row=r_end, end_column=20)
                ws.cell(row=r_start, column=20, value=prev)

                ws.merge_cells(start_row=r_start, start_column=21, end_row=r_end, end_column=21)
                ws.cell(row=r_start, column=21, value=sit_txt)

                for idx, e in enumerate([1, 2, 3]):
                    r = r_start + idx
                    s_d = dias_map.loc["Segunda", f"Etapa {e}"]
                    t_d = dias_map.loc["Terça", f"Etapa {e}"]
                    q_d = dias_map.loc["Quarta", f"Etapa {e}"]
                    qui_d = dias_map.loc["Quinta", f"Etapa {e}"]
                    sex_d = dias_map.loc["Sexta", f"Etapa {e}"]

                    s_t, t_t, q_t, qui_t, sex_t = s_a * s_d, t_a * t_d, q_a * q_d, qui_a * qui_d, sex_a * sex_d
                    tot_etapa = s_t + t_t + q_t + qui_t + sex_t

                    ws.cell(row=r, column=2, value=e)
                    ws.cell(row=r, column=4, value=s_d)
                    ws.cell(row=r, column=5, value=s_t)
                    ws.cell(row=r, column=7, value=t_d)
                    ws.cell(row=r, column=8, value=t_t)
                    ws.cell(row=r, column=10, value=q_d)
                    ws.cell(row=r, column=11, value=q_t)
                    ws.cell(row=r, column=13, value=qui_d)
                    ws.cell(row=r, column=14, value=qui_t)
                    ws.cell(row=r, column=16, value=sex_d)
                    ws.cell(row=r, column=17, value=sex_t)
                    ws.cell(row=r, column=18, value=tot_etapa)

                # Formatação de bordas e estilos do bloco
                for r in range(r_start, r_end + 1):
                    for c in range(1, 22):
                        cell = ws.cell(row=r, column=c)
                        cell.alignment = align_center
                        cell.border = border_cell
                        cell.fill = bg_fill
                        cell.font = font_regular
                        # Destaques especiais
                        if c in [3, 6, 9, 12, 15] and cell.value and cell.value > 0:
                            cell.fill = destaque_fill
                            cell.font = font_bold
                        elif c in [1, 18, 19, 20]:
                            cell.font = font_bold
                        if c == 21:
                            cell.fill = sit_fill
                            cell.font = sit_font

                current_row += 3

            # Larguras das colunas
            ws.column_dimensions["A"].width = 24
            ws.column_dimensions["B"].width = 8
            for col in ["C", "D", "E", "F", "G", "H", "I", "J", "K", "L", "M", "N", "O", "P", "Q"]:
                ws.column_dimensions[col].width = 11
            ws.column_dimensions["R"].width = 16
            ws.column_dimensions["S"].width = 14
            ws.column_dimensions["T"].width = 14
            ws.column_dimensions["U"].width = 16

            excel_buffer = io.BytesIO()
            wb.save(excel_buffer)
            st.download_button(
                label="📥 Baixar Planilha Padrão (.xlsx)",
                data=excel_buffer.getvalue(),
                file_name=f"Conferencia_{st.session_state.turma_atual}_{st.session_state.ano_letivo}.xlsx",
                mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
    else:
        st.info(
            f"Nenhuma disciplina lançada para a Turma {st.session_state.turma_atual}."
        )
